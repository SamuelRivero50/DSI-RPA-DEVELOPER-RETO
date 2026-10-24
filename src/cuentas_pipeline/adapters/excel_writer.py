import os
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from cuentas_pipeline.core.config import SalidaConfig
from cuentas_pipeline.core.models import Estado

_COLUMNAS = [
    "Línea origen",
    "Valor original",
    "ID cuenta",
    "Estado",
    "Causa",
    "Apariciones",
    "Intentos",
    "Detalle del error",
    "Consultado en",
    "ID usuario",
    "Título",
    "Cuerpo",
]

_PREFIJOS_FORMULA = ("=", "+", "-", "@")

_FUENTE = "Calibri"

_RELLENOS = {
    Estado.OK.value: "C6EFCE",
    Estado.NO_ENCONTRADO.value: "FFEB9C",
    Estado.DATO_INVALIDO.value: "F8CBAD",
    Estado.ERROR_TECNICO.value: "FFC7CE",
    Estado.PENDIENTE.value: "D9D9D9",
}

_RELLENOS_SUAVES = {
    Estado.NO_ENCONTRADO.value: "FFF2CC",
    Estado.DATO_INVALIDO.value: "FCE4D6",
    Estado.ERROR_TECNICO.value: "FFE6E6",
    Estado.PENDIENTE.value: "F2F2F2",
}

_COLOR_BANDA = "404040"
_COLOR_TEXTO_SECUNDARIO = "808080"
_LADO_FINO = Side(style="thin", color="BFBFBF")
_BORDE_FINO = Border(left=_LADO_FINO, right=_LADO_FINO, top=_LADO_FINO, bottom=_LADO_FINO)
_BORDE_TOTAL = Border(
    left=_LADO_FINO,
    right=_LADO_FINO,
    top=Side(style="thin", color=_COLOR_BANDA),
    bottom=_LADO_FINO,
)

_ANCHOS = [12, 16, 11, 18, 22, 13, 10, 24, 20, 12, 40, 70]

_COLUMNAS_CENTRADAS = {1, 3, 6, 7, 10}
_ALINEACION_CENTRO = Alignment(horizontal="center", vertical="center")
_ALINEACION_IZQUIERDA = Alignment(horizontal="left", vertical="center")

_ORDEN_ESTADOS = (
    Estado.OK,
    Estado.NO_ENCONTRADO,
    Estado.DATO_INVALIDO,
    Estado.ERROR_TECNICO,
    Estado.PENDIENTE,
)


def _fecha_archivo(filas: list[dict[str, object]]) -> str:
    fechas = [str(fila["creado_en"]) for fila in filas if fila.get("creado_en")]
    if not fechas:
        return datetime.now().astimezone().date().isoformat()
    return min(fechas)[:10]


def _valores(fila: dict[str, object]) -> list[object]:
    estado = str(fila["estado"])
    id_cuenta = None if estado == Estado.DATO_INVALIDO.value else int(fila["id"])
    return [
        fila["linea_origen"],
        fila["valor_original"],
        id_cuenta,
        estado,
        fila["causa"],
        fila["apariciones"],
        fila["intentos"],
        fila["detalle_error"],
        _formatear_fecha(str(fila["actualizado_en"])),
        fila["user_id"],
        fila["titulo"],
        fila["cuerpo"],
    ]


def _fuente(
    *,
    tamano: int = 11,
    negrita: bool = False,
    cursiva: bool = False,
    color: str | None = None,
) -> Font:
    return Font(name=_FUENTE, size=tamano, bold=negrita, italic=cursiva, color=color)


def _relleno(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def _proteger_texto(celda) -> None:
    if isinstance(celda.value, str) and celda.value.startswith(_PREFIJOS_FORMULA):
        celda.data_type = "s"


def _escribir_celda(
    hoja,
    fila: int,
    columna: int,
    valor,
    *,
    tamano: int = 11,
    negrita: bool = False,
    cursiva: bool = False,
    color: str | None = None,
    relleno: str | None = None,
    borde=None,
):
    celda = hoja.cell(row=fila, column=columna, value=valor)
    celda.font = _fuente(tamano=tamano, negrita=negrita, cursiva=cursiva, color=color)
    if relleno is not None:
        celda.fill = _relleno(relleno)
    if borde is not None:
        celda.border = borde
    _proteger_texto(celda)
    return celda


def _primera_carga(filas: list[dict[str, object]]) -> str:
    fechas = [str(fila["creado_en"]) for fila in filas if fila.get("creado_en")]
    return min(fechas) if fechas else ""


def _formatear_fecha(valor: str) -> str:
    if not valor:
        return ""
    try:
        return datetime.fromisoformat(valor).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return valor


def _aplicar_alineacion(celda, indice: int) -> None:
    if indice in _COLUMNAS_CENTRADAS:
        celda.alignment = _ALINEACION_CENTRO
    else:
        celda.alignment = _ALINEACION_IZQUIERDA


def _conteos_por_estado(filas: list[dict[str, object]]) -> dict[str, int]:
    conteos = {estado.value: 0 for estado in Estado}
    for fila in filas:
        conteos[str(fila["estado"])] += 1
    return conteos


def _conteos_por_causa(filas: list[dict[str, object]]) -> dict[str, int]:
    conteos: dict[str, int] = {}
    for fila in filas:
        if fila["estado"] != Estado.ERROR_TECNICO.value or not fila["causa"]:
            continue
        causa = str(fila["causa"])
        conteos[causa] = conteos.get(causa, 0) + 1
    return dict(sorted(conteos.items()))


def _banda(hoja, fila: int, titulo: str) -> None:
    for columna in (1, 2):
        _escribir_celda(
            hoja,
            fila,
            columna,
            titulo if columna == 1 else None,
            negrita=True,
            color="FFFFFF",
            relleno=_COLOR_BANDA,
        )
    hoja.row_dimensions[fila].height = 20


def _cabecera_tabla(hoja, fila: int, *titulos: str) -> None:
    for columna, titulo in enumerate(titulos, start=1):
        _escribir_celda(
            hoja,
            fila,
            columna,
            titulo,
            negrita=True,
            relleno=_RELLENOS[Estado.PENDIENTE.value],
            borde=_BORDE_FINO,
        )


def _campo(hoja, fila: int, etiqueta: str, valor, **estilo) -> None:
    _escribir_celda(hoja, fila, 1, etiqueta)
    _escribir_celda(hoja, fila, 2, valor, **estilo)


def _escribir_resumen(
    libro,
    archivo_nombre: str,
    archivo_hash: str,
    normalizacion,
    filas: list[dict[str, object]],
    generado_en: str,
) -> None:
    hoja = libro.create_sheet("Resumen")
    hoja.sheet_view.showGridLines = False
    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 16

    _escribir_celda(hoja, 1, 1, "Resumen del archivo", tamano=14, negrita=True)
    hoja.row_dimensions[1].height = 22
    fila = 3

    _banda(hoja, fila, "Archivo")
    fila += 1
    _campo(hoja, fila, "Nombre", archivo_nombre)
    fila += 1
    _campo(
        hoja,
        fila,
        "Hash",
        archivo_hash,
        tamano=9,
        color=_COLOR_TEXTO_SECUNDARIO,
    )
    fila += 1
    _campo(hoja, fila, "Primera carga", _formatear_fecha(_primera_carga(filas)))
    fila += 1
    _campo(hoja, fila, "Generado en", _formatear_fecha(generado_en))
    fila += 2

    _banda(hoja, fila, "Entrada")
    fila += 1
    for etiqueta, valor in [
        ("Líneas con datos", normalizacion.lineas_con_datos),
        ("Líneas en blanco ignoradas", normalizacion.lineas_en_blanco),
        ("Repeticiones ignoradas", normalizacion.repeticiones_ignoradas),
        ("Identificadores únicos", len(normalizacion.filas)),
    ]:
        _campo(hoja, fila, etiqueta, valor)
        fila += 1
    fila += 1

    _banda(hoja, fila, "Resultado por estado")
    fila += 1
    _cabecera_tabla(hoja, fila, "Estado", "Total")
    fila += 1
    conteos = _conteos_por_estado(filas)
    for estado in _ORDEN_ESTADOS:
        total = conteos[estado.value]
        for columna, valor in [(1, estado.value), (2, total)]:
            _escribir_celda(
                hoja,
                fila,
                columna,
                valor,
                relleno=_RELLENOS[estado.value],
                borde=_BORDE_FINO,
            )
        fila += 1
    for columna, valor in [(1, "Total"), (2, sum(conteos.values()))]:
        _escribir_celda(
            hoja,
            fila,
            columna,
            valor,
            negrita=True,
            borde=_BORDE_TOTAL,
        )
    fila += 2

    _banda(hoja, fila, "Errores técnicos por causa")
    fila += 1
    errores = _conteos_por_causa(filas)
    if errores:
        _cabecera_tabla(hoja, fila, "Causa", "Total")
        fila += 1
        for causa, total in errores.items():
            for columna, valor in [(1, causa), (2, total)]:
                _escribir_celda(hoja, fila, columna, valor, borde=_BORDE_FINO)
            fila += 1
    else:
        _escribir_celda(
            hoja,
            fila,
            1,
            "Sin errores técnicos",
            cursiva=True,
            color=_COLOR_TEXTO_SECUNDARIO,
        )
        fila += 1
    fila += 1

    _banda(hoja, fila, "Repetidos")
    fila += 1
    _cabecera_tabla(hoja, fila, "Id", "Apariciones")
    fila += 1
    if normalizacion.repetidos:
        for id_cuenta, apariciones in normalizacion.repetidos:
            valor_id = int(id_cuenta) if id_cuenta.isdigit() else id_cuenta
            for columna, valor in [(1, valor_id), (2, apariciones)]:
                _escribir_celda(hoja, fila, columna, valor, borde=_BORDE_FINO)
            fila += 1
    else:
        _escribir_celda(
            hoja,
            fila,
            1,
            "Sin repetidos",
            cursiva=True,
            color=_COLOR_TEXTO_SECUNDARIO,
        )


class ExcelWriter:
    def __init__(self, config: SalidaConfig):
        self._carpeta = Path(config.carpeta)

    def escribir(
        self,
        archivo_nombre: str,
        archivo_hash: str,
        normalizacion,
        filas: list[dict[str, object]],
    ) -> Path:
        self._carpeta.mkdir(parents=True, exist_ok=True)
        fecha = _fecha_archivo(filas)
        ruta = self._carpeta / f"consolidado_{fecha}_{archivo_hash[:8]}.xlsx"

        descriptor, nombre_temporal = tempfile.mkstemp(
            prefix=".consolidado_",
            suffix=".xlsx",
            dir=self._carpeta,
        )
        os.close(descriptor)
        temporal = Path(nombre_temporal)

        try:
            libro = Workbook()
            hoja = libro.active
            hoja.title = "Resultados"
            hoja.append(_COLUMNAS)
            hoja.row_dimensions[1].height = 28
            for celda in hoja[1]:
                celda.font = _fuente(negrita=True, color="FFFFFF")
                celda.fill = _relleno(_COLOR_BANDA)
                celda.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )

            for fila in filas:
                estado = str(fila["estado"])
                hoja.append(_valores(fila))
                hoja.row_dimensions[hoja.max_row].height = 20
                relleno_suave = _RELLENOS_SUAVES.get(estado)
                for indice, celda in enumerate(hoja[hoja.max_row], start=1):
                    _proteger_texto(celda)
                    _aplicar_alineacion(celda, indice)
                    if relleno_suave is not None:
                        celda.fill = _relleno(relleno_suave)
                hoja.cell(row=hoja.max_row, column=4).fill = _relleno(
                    _RELLENOS[estado]
                )

            for indice, ancho in enumerate(_ANCHOS, start=1):
                hoja.column_dimensions[get_column_letter(indice)].width = ancho

            hoja.freeze_panes = "D2"
            hoja.auto_filter.ref = f"A1:L{hoja.max_row}"
            hoja.page_setup.orientation = "landscape"
            hoja.page_setup.fitToWidth = 1
            hoja.page_setup.fitToHeight = 0
            hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            hoja.print_title_rows = "1:1"

            _escribir_resumen(
                libro,
                archivo_nombre,
                archivo_hash,
                normalizacion,
                filas,
                datetime.now().astimezone().isoformat(timespec="seconds"),
            )

            libro.active = 0
            libro.save(temporal)
            os.replace(temporal, ruta)
        finally:
            if temporal.exists():
                temporal.unlink()

        return ruta
