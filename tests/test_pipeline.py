from dataclasses import replace
from pathlib import Path

import pytest
import requests
from openpyxl import load_workbook

from cuentas_pipeline.adapters.csv_reader import LecturaCsv, leer_csv
from cuentas_pipeline.adapters.excel_writer import ExcelWriter
from cuentas_pipeline.adapters.repository import Repositorio
from cuentas_pipeline.core.config import (
    ApiConfig,
    Config,
    EntradaConfig,
    ReintentosConfig,
    SalidaConfig,
)
from cuentas_pipeline.core.models import Estado
from cuentas_pipeline.pipeline import ejecutar


class RespuestaFalsa:
    def __init__(self, status_code, datos=None, headers=None):
        self.status_code = status_code
        self._datos = datos
        self.headers = headers or {}

    def json(self):
        return self._datos


class ClienteFalso:
    def __init__(self, responder):
        self._responder = responder
        self.llamadas = []

    def obtener(self, id_cuenta):
        self.llamadas.append(id_cuenta)
        return self._responder(id_cuenta)


def _lectura(filas, archivo_hash="a" * 64, nombre="cuentas.csv"):
    return LecturaCsv(nombre, archivo_hash, archivo_hash[:8], filas)


def _config(tmp_path, max_intentos=4):
    return Config(
        entrada=EntradaConfig(),
        reintentos=ReintentosConfig(max_intentos=max_intentos),
        salida=SalidaConfig(
            carpeta=str(tmp_path / "salida"),
            base_datos=str(tmp_path / "estado" / "estado.db"),
        ),
    )


def _respuesta_ok(id_cuenta):
    valor = int(id_cuenta)
    return RespuestaFalsa(
        200,
        {"userId": 1, "id": valor, "title": f"titulo {valor}", "body": f"cuerpo {valor}"},
    )


def test_el_csv_del_reto_termina_con_los_conteos_esperados(tmp_path):
    config = _config(tmp_path)
    lectura = leer_csv("cuentas.csv", EntradaConfig())

    def responder(id_cuenta):
        if id_cuenta in {"0", "101", "150"}:
            return RespuestaFalsa(404)
        return _respuesta_ok(id_cuenta)

    cliente = ClienteFalso(responder)
    repo = Repositorio(config.salida)

    resumen = ejecutar(lectura, config, repo, cliente, sleep=lambda _segundos: None)

    assert len(cliente.llamadas) == 16
    assert resumen.por_estado == {
        Estado.PENDIENTE: 0,
        Estado.OK: 13,
        Estado.NO_ENCONTRADO: 3,
        Estado.DATO_INVALIDO: 1,
        Estado.ERROR_TECNICO: 0,
    }
    assert resumen.lineas_con_datos == 20
    assert resumen.lineas_en_blanco == 1
    assert resumen.repeticiones_ignoradas == 3
    assert resumen.ids_unicos == 17
    assert resumen.repetidos == [("7", 3), ("12", 2)]
    assert resumen.consultados == 16
    assert resumen.reutilizados == 0
    assert resumen.reintentos_totales == 0
    assert resumen.duracion_s >= 0
    assert repo.pendientes(lectura.archivo_hash) == []
    repo.cerrar()


def test_la_segunda_ejecucion_no_hace_llamadas(tmp_path):
    config = _config(tmp_path)
    lectura = leer_csv("cuentas.csv", EntradaConfig())

    def responder(id_cuenta):
        if id_cuenta in {"0", "101", "150"}:
            return RespuestaFalsa(404)
        return _respuesta_ok(id_cuenta)

    cliente = ClienteFalso(responder)
    repo = Repositorio(config.salida)
    ejecutar(lectura, config, repo, cliente, sleep=lambda _segundos: None)
    llamadas_primera = len(cliente.llamadas)

    resumen = ejecutar(lectura, config, repo, cliente, sleep=lambda _segundos: None)

    assert len(cliente.llamadas) == llamadas_primera
    assert resumen.consultados == 0
    assert resumen.reutilizados == 16
    assert resumen.por_estado[Estado.OK] == 13
    assert resumen.por_estado[Estado.NO_ENCONTRADO] == 3
    assert resumen.por_estado[Estado.DATO_INVALIDO] == 1
    repo.cerrar()


def test_un_error_tecnico_se_reintenta_al_relanzar(tmp_path):
    config = _config(tmp_path, max_intentos=2)
    lectura = _lectura([(2, "7")])
    repo = Repositorio(config.salida)

    def sin_red(_id_cuenta):
        raise requests.exceptions.ConnectionError("sin red")

    primer_cliente = ClienteFalso(sin_red)
    primer_resumen = ejecutar(
        lectura, config, repo, primer_cliente, sleep=lambda _segundos: None
    )

    assert primer_resumen.por_estado[Estado.ERROR_TECNICO] == 1
    assert primer_resumen.consultados == 1
    assert primer_resumen.reintentos_totales == 1
    assert [fila.id for fila in repo.pendientes(lectura.archivo_hash)] == ["7"]

    segundo_cliente = ClienteFalso(_respuesta_ok)
    segundo_resumen = ejecutar(
        lectura, config, repo, segundo_cliente, sleep=lambda _segundos: None
    )

    assert segundo_cliente.llamadas == ["7"]
    assert segundo_resumen.consultados == 1
    assert segundo_resumen.reutilizados == 0
    assert segundo_resumen.por_estado[Estado.OK] == 1
    repo.cerrar()


def test_una_excepcion_inesperada_no_detiene_los_demas_ids(tmp_path):
    config = _config(tmp_path)
    lectura = _lectura([(2, "7"), (3, "8"), (4, "9")])
    repo = Repositorio(config.salida)

    def responder(id_cuenta):
        if id_cuenta == "8":
            raise ValueError("bug")
        return _respuesta_ok(id_cuenta)

    cliente = ClienteFalso(responder)

    resumen = ejecutar(lectura, config, repo, cliente, sleep=lambda _segundos: None)

    assert cliente.llamadas == ["7", "8", "9"]
    assert resumen.consultados == 3
    filas = {fila["id"]: fila for fila in repo.filas(lectura.archivo_hash)}
    assert filas["7"]["estado"] == "OK"
    assert filas["8"]["estado"] == "ERROR_TECNICO"
    assert filas["8"]["causa"] == "ERROR_INESPERADO"
    assert filas["9"]["estado"] == "OK"
    repo.cerrar()


def test_un_dato_invalido_no_se_consulta(tmp_path):
    config = _config(tmp_path)
    lectura = _lectura([(2, "7"), (3, "abc")])
    repo = Repositorio(config.salida)
    cliente = ClienteFalso(_respuesta_ok)

    resumen = ejecutar(lectura, config, repo, cliente, sleep=lambda _segundos: None)

    assert cliente.llamadas == ["7"]
    assert resumen.consultados == 1
    assert resumen.por_estado[Estado.DATO_INVALIDO] == 1
    assert resumen.por_estado[Estado.OK] == 1
    repo.cerrar()


def test_un_csv_solo_con_encabezado_no_consulta_nada(tmp_path):
    config = _config(tmp_path)
    lectura = _lectura([])
    repo = Repositorio(config.salida)
    cliente = ClienteFalso(_respuesta_ok)

    resumen = ejecutar(lectura, config, repo, cliente, sleep=lambda _segundos: None)

    assert cliente.llamadas == []
    assert resumen.consultados == 0
    assert resumen.ids_unicos == 0
    assert resumen.por_estado == {estado: 0 for estado in Estado}
    repo.cerrar()


def test_la_pausa_entre_consultas_solo_ocurre_entre_ids(tmp_path):
    config = replace(
        _config(tmp_path),
        api=ApiConfig(pausa_entre_consultas_s=0.5),
        reintentos=ReintentosConfig(max_intentos=1),
    )
    lectura = _lectura([(2, "7"), (3, "8"), (4, "9")])
    repo = Repositorio(config.salida)
    cliente = ClienteFalso(_respuesta_ok)
    esperas = []

    ejecutar(lectura, config, repo, cliente, sleep=esperas.append)

    assert esperas == [0.5, 0.5]
    repo.cerrar()


def test_si_el_escritor_falla_la_excepcion_se_propaga(tmp_path):
    config = _config(tmp_path)
    lectura = _lectura([(2, "7")])
    repo = Repositorio(config.salida)

    class EscritorRoto:
        def escribir(self, *_args):
            raise OSError("sin disco")

    with pytest.raises(OSError, match="sin disco"):
        ejecutar(
            lectura,
            config,
            repo,
            ClienteFalso(_respuesta_ok),
            escritor=EscritorRoto(),
            sleep=lambda _segundos: None,
        )

    repo.cerrar()


def test_el_excel_tiene_una_fila_por_id_unico(tmp_path):
    config = _config(tmp_path)
    lectura = leer_csv("cuentas.csv", EntradaConfig())

    def responder(id_cuenta):
        if id_cuenta in {"0", "101", "150"}:
            return RespuestaFalsa(404)
        return _respuesta_ok(id_cuenta)

    repo = Repositorio(config.salida)
    resumen = ejecutar(
        lectura,
        config,
        repo,
        ClienteFalso(responder),
        escritor=ExcelWriter(config.salida),
        sleep=lambda _segundos: None,
    )

    assert resumen.ruta_excel is not None
    assert resumen.ruta_excel.exists()
    assert resumen.ruta_excel.name.startswith("consolidado_")
    assert resumen.ruta_excel.name.endswith(f"_{lectura.hash8}.xlsx")
    assert list(Path(config.salida.carpeta).glob(".consolidado_*")) == []

    hoja = load_workbook(resumen.ruta_excel)["Resultados"]
    assert hoja.max_row == 18
    assert hoja.max_column == 12
    assert [celda.value for celda in hoja[1]] == [
        "linea_origen",
        "valor_original",
        "id_cuenta",
        "estado",
        "causa",
        "apariciones",
        "user_id",
        "titulo",
        "cuerpo",
        "intentos",
        "detalle_error",
        "consultado_en",
    ]

    por_id = {fila[2].value: fila for fila in hoja.iter_rows(min_row=2)}
    assert len(por_id) == 17
    assert por_id[7][1].value == "7"
    assert por_id[7][3].value == "OK"
    assert por_id[7][5].value == 3
    assert por_id[7][7].value == "titulo 7"
    assert por_id[23][1].value == " 23"
    assert por_id[101][3].value == "NO_ENCONTRADO"
    assert por_id[101][4].value == "HTTP_404"
    assert por_id[None][1].value == "abc"
    assert por_id[None][2].value is None
    assert por_id[None][3].value == "DATO_INVALIDO"
    repo.cerrar()


def test_el_texto_con_prefijo_de_formula_se_escribe_como_texto(tmp_path):
    config = _config(tmp_path)
    lectura = _lectura([(2, "7"), (3, "=cmd")])

    def responder(id_cuenta):
        return RespuestaFalsa(200, {"id": 7, "title": "=2+2", "body": "@hola"})

    repo = Repositorio(config.salida)
    resumen = ejecutar(
        lectura,
        config,
        repo,
        ClienteFalso(responder),
        escritor=ExcelWriter(config.salida),
        sleep=lambda _segundos: None,
    )

    hoja = load_workbook(resumen.ruta_excel)["Resultados"]
    por_id = {fila[2].value: fila for fila in hoja.iter_rows(min_row=2)}
    assert por_id[7][7].value == "=2+2"
    assert por_id[7][7].data_type == "s"
    assert por_id[7][8].value == "@hola"
    assert por_id[7][8].data_type == "s"
    assert por_id[None][1].value == "=cmd"
    assert por_id[None][1].data_type == "s"
    repo.cerrar()


def test_el_excel_tiene_formato_y_color_por_estado(tmp_path):
    config = _config(tmp_path)
    lectura = _lectura([(2, "7"), (3, "101"), (4, "8"), (5, "abc")])

    def responder(id_cuenta):
        if id_cuenta == "101":
            return RespuestaFalsa(404)
        if id_cuenta == "8":
            raise ValueError("bug")
        return RespuestaFalsa(200, {"id": 7, "title": "titulo", "body": "cuerpo"})

    repo = Repositorio(config.salida)
    resumen = ejecutar(
        lectura,
        config,
        repo,
        ClienteFalso(responder),
        escritor=ExcelWriter(config.salida),
        sleep=lambda _segundos: None,
    )

    hoja = load_workbook(resumen.ruta_excel)["Resultados"]
    assert hoja.freeze_panes == "A2"
    assert hoja.auto_filter.ref == f"A1:L{hoja.max_row}"
    assert hoja["A1"].font.bold
    por_id = {fila[2].value: fila for fila in hoja.iter_rows(min_row=2)}
    assert por_id[7][8].alignment.wrap_text
    assert por_id[7][0].fill.fgColor.rgb.endswith("C6EFCE")
    assert por_id[101][0].fill.fgColor.rgb.endswith("FFEB9C")
    assert por_id[8][0].fill.fgColor.rgb.endswith("FFC7CE")
    assert por_id[None][0].fill.fgColor.rgb.endswith("F8CBAD")
    repo.cerrar()
